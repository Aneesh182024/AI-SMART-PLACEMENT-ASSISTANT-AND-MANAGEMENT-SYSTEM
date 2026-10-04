import io
import csv
import openpyxl
from openpyxl.styles import PatternFill

from app.database_mock import STUDENTS_DB
from app.excel_service import ExcelService

def run_excel_service_tests():
    print("=" * 70)
    print("PSNA IT PLACEMENT ASSISTANT: EXCEL SERVICE VALIDATION SUITE")
    print("=" * 70)

    # --------------------------------------------------------------------------
    # TEST 1: SMART BULK EXCEL INGESTION (.CSV & .XLSX) WITH DE-DUPLICATION
    # --------------------------------------------------------------------------
    print("\n[TEST 1] Testing Smart Bulk Ingestion with De-Duplication...")
    
    # Check initial record for student 713821104001
    initial_name = STUDENTS_DB["713821104001"]["name"]
    initial_cgpa = STUDENTS_DB["713821104001"]["cgpa"]
    print(f"  Before Ingest: 713821104001 -> Name: '{initial_name}', CGPA: {initial_cgpa}")

    # Create in-memory CSV with:
    # 1 existing student (with updated CGPA, mobile, but a modified name to verify it keeps canonical name or updates cleanly without duplication)
    # 1 new student
    csv_content = """Register No,Name,Gender,Year,Section,Email,Mobile,CGPA,Arrears
713821104001,Aneesh Kanna N (Duplicate Name Attempt),Male,IV,A,aneesh.new@psnacet.edu.in,9998887771,9.85,0
713821104099,Karthika Devi S,Female,IV,C,karthika@psnacet.edu.in,9876543299,8.72,0
"""
    csv_bytes = csv_content.encode("utf-8")
    result_csv = ExcelService.ingest_spreadsheet(csv_bytes, "students_roster_batch.csv")
    print(f"  CSV Ingest Result: {result_csv}")

    assert result_csv["status"] == "success"
    assert result_csv["updated_count"] == 1
    assert result_csv["inserted_count"] == 1

    # Verify student 713821104001: name preserved without duplication, CGPA updated to 9.85
    assert STUDENTS_DB["713821104001"]["name"] == initial_name, "Student name should not be duplicated/corrupted on upsert!"
    assert STUDENTS_DB["713821104001"]["cgpa"] == 9.85, "CGPA was not updated properly!"
    assert STUDENTS_DB["713821104001"]["mobile"] == "9998887771"

    # Verify new student 713821104099 inserted
    assert "713821104099" in STUDENTS_DB
    assert STUDENTS_DB["713821104099"]["name"] == "Karthika Devi S"
    assert STUDENTS_DB["713821104099"]["gender"] == "Female"
    assert STUDENTS_DB["713821104099"]["section"] == "C"
    print("  [SUCCESS] Smart Upsert verified: de-duplicated existing students and inserted new roster records.")

    # --------------------------------------------------------------------------
    # TEST 2: TEST XLSX INGESTION
    # --------------------------------------------------------------------------
    print("\n[TEST 2] Testing .XLSX Ingestion...")
    wb_test = openpyxl.Workbook()
    ws_test = wb_test.active
    ws_test.append(["Register No", "Name", "Gender", "Year", "Section", "Email", "Mobile", "CGPA", "Arrears"])
    ws_test.append(["713821104004", "Dharani S", "Female", "IV", "B", "dharani@psnacet.edu.in", "9876543213", "8.10", "1"])
    ws_test.append(["713821104100", "Vigneshwaran P", "Male", "IV", "D", "vignesh@psnacet.edu.in", "9876543290", "7.95", "0"])
    
    xlsx_stream = io.BytesIO()
    wb_test.save(xlsx_stream)
    xlsx_bytes = xlsx_stream.getvalue()

    result_xlsx = ExcelService.ingest_spreadsheet(xlsx_bytes, "annual_roster.xlsx")
    print(f"  XLSX Ingest Result: {result_xlsx}")
    assert result_xlsx["status"] == "success"
    assert "713821104100" in STUDENTS_DB
    print("  [SUCCESS] .XLSX binary ingestion parsed successfully.")

    # --------------------------------------------------------------------------
    # TEST 3: 5-TIER COLOR-CODED EXCEL EXPORT ENGINE
    # --------------------------------------------------------------------------
    print("\n[TEST 3] Testing 5-Tier Color-Coded Excel Generation...")
    stream = ExcelService.generate_5tier_color_roster()
    exported_bytes = stream.getvalue()
    print(f"  Generated workbook size: {len(exported_bytes)} bytes")
    assert len(exported_bytes) > 5000, "Workbook is too small!"

    # Load exported workbook to inspect styling and headers
    wb_exported = openpyxl.load_workbook(io.BytesIO(exported_bytes), data_only=False)
    ws_exp = wb_exported.active

    # Check Header Banner
    col_a1 = ws_exp["A1"].value
    col_a2 = ws_exp["A2"].value
    col_a3 = ws_exp["A3"].value
    print(f"  Row 1 Banner: '{col_a1}'")
    print(f"  Row 2 Subtitle: '{col_a2}'")
    print(f"  Row 3 Developers: '{col_a3}'")
    assert "PSNA COLLEGE OF ENGINEERING AND TECHNOLOGY" in col_a1
    assert "DEPARTMENT OF INFORMATION TECHNOLOGY" in col_a2
    assert "ANEESH KANNA N and ANNE BENILDA A" in col_a3

    # Check Table Headers (Row 5)
    headers_row = [cell.value for cell in ws_exp[5] if cell.value]
    print(f"  Headers ({len(headers_row)}): {headers_row}")
    assert "Register No" in headers_row
    assert "Subject Code & Title" in headers_row
    assert "Semester" in headers_row

    # Verify 5-Tier Color Coding across rows
    print("\n  Inspecting Cell Fills for 5-Tier Color Validation:")
    tier_found = {1: False, 2: False, 3: False, 4: False, 5: False}

    for row_idx in range(6, ws_exp.max_row + 1):
        reg_cell = ws_exp.cell(row=row_idx, column=1)
        sub_cell = ws_exp.cell(row=row_idx, column=10)
        sem_cell = ws_exp.cell(row=row_idx, column=11)

        reg_color = reg_cell.fill.start_color.rgb if reg_cell.fill and reg_cell.fill.start_color else None
        sub_color = sub_cell.fill.start_color.rgb if sub_cell.fill and sub_cell.fill.start_color else None
        sem_color = sem_cell.fill.start_color.rgb if sem_cell.fill and sem_cell.fill.start_color else None

        # Check Tier 5: Bright Gold (#EAB308) on Register No (High CGPA Top Performer)
        if reg_color and "EAB308" in str(reg_color).upper():
            tier_found[5] = True
            print(f"    [TIER 5 MATCH] Row {row_idx} ({reg_cell.value}): Bright Gold (#EAB308) Top Performer")

        # Check Tier 1: Emerald Green (#10B981) on Register No (Without Arrear)
        elif reg_color and "10B981" in str(reg_color).upper():
            tier_found[1] = True
            print(f"    [TIER 1 MATCH] Row {row_idx} ({reg_cell.value}): Emerald Green (#10B981) Without Arrear")

        # Check Tier 2: Crimson Red (#EF4444) on Specific Failed Subject (Current Arrear)
        if sub_color and "EF4444" in str(sub_color).upper():
            tier_found[2] = True
            print(f"    [TIER 2 MATCH] Row {row_idx} ({sub_cell.value}): Crimson Red (#EF4444) Current Arrear")

        # Check Tier 4: Soft Pastel Blue (#3B82F6) on Cleared Subject + Cleared Sem (Previous Arrear Cleared)
        if sub_color and "3B82F6" in str(sub_color).upper() and sem_color and "3B82F6" in str(sem_color).upper():
            tier_found[4] = True
            print(f"    [TIER 4 MATCH] Row {row_idx} ({sub_cell.value} | {sem_cell.value}): Pastel Blue (#3B82F6) Cleared Arrear")

        # Check Tier 3: Amber Orange (#F59E0B) on Subject + Semester (Previous Arrear)
        if sub_color and "F59E0B" in str(sub_color).upper() and sem_color and "F59E0B" in str(sem_color).upper():
            tier_found[3] = True
            print(f"    [TIER 3 MATCH] Row {row_idx} ({sub_cell.value} | {sem_cell.value}): Amber Orange (#F59E0B) Previous Arrear")

    print(f"\n  Tiers Identified: {tier_found}")
    for t, found in tier_found.items():
        assert found, f"Tier {t} color pattern was not triggered!"

    print("\n[ALL EXCEL INGESTION & 5-TIER COLOR EXPORT TESTS PASSED SUCCESSFULLY!]")
    print("=" * 70)

if __name__ == "__main__":
    run_excel_service_tests()
